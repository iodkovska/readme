"""Filling the fund's Word template with a completed scoring memo.

The template ships with a worked example in it (Armeta). Every field is
overwritten, so nothing from that example can leak into a generated memo — the
tests assert exactly that.

Layout of `templates/scoring_template.docx`, by table index:
    0   13x4 deal facts — col 1 holds the left column's values, col 3 the right's
    1   2x10 scores     — row 1 is the ten category scores, in template order
    2-10 2x1 sections   — row 0 is the heading, row 1 the content
"""

from __future__ import annotations

import io
from pathlib import Path

from docx import Document
from docx.shared import Pt

from .scoring import (
    CATEGORIES,
    FACT_LABELS_LEFT,
    FACT_LABELS_RIGHT,
    MAX_TOTAL,
    ScoringMemo,
)

TEMPLATE_PATH = Path(__file__).parent / "templates" / "scoring_template.docx"

# Table 0 rows the bot fills. The left column's values go in cell 1, the right
# column's in cell 3; rows not listed here (VP, GP, Contacts, and the two dates)
# are the fund's own workflow data and are deliberately left blank.
LEFT_ROWS = list(enumerate(FACT_LABELS_LEFT))        # rows 0-9
RIGHT_ROWS = list(enumerate(FACT_LABELS_RIGHT, 2))   # rows 2-11
SCORING_ROW = 12                                     # right column: "Scoring"
BLANK_LEFT_ROWS = (10, 11, 12)                       # Contacts, Commitment, Money Transfer
BLANK_RIGHT_ROWS = (0, 1)                            # VP, GP

SECTION_TABLES = {
    2: "product",
    3: "business_model",
    4: "traction",
    5: "team",
    6: "go_to_market",
    7: "market",
    8: "competitors",
    9: "technology",
    10: "deal",
}


def _set_cell(cell, text: str) -> None:
    """Replace a cell's text, keeping the formatting of its first run."""
    text = (text or "").strip()
    paragraph = cell.paragraphs[0]

    # Drop every paragraph after the first, then every run after the first.
    for extra in cell.paragraphs[1:]:
        extra._element.getparent().remove(extra._element)
    for run in paragraph.runs[1:]:
        run._element.getparent().remove(run._element)

    if paragraph.runs:
        paragraph.runs[0].text = text
    elif text:
        paragraph.add_run(text)


def _fill_cell_lines(cell, lines: list[tuple[str, str]]) -> None:
    """Write `label: value` lines into a section cell, label in bold."""
    paragraph = cell.paragraphs[0]
    for extra in cell.paragraphs[1:]:
        extra._element.getparent().remove(extra._element)
    for run in list(paragraph.runs):
        run._element.getparent().remove(run._element)

    first = True
    for label, value in lines:
        target = paragraph if first else cell.add_paragraph()
        first = False
        if label:
            run = target.add_run(f"{label}: ")
            run.bold = True
        body = target.add_run((value or "—").strip())
        body.font.size = Pt(10)


def _delete_between(doc, start_text: str, end_text: str) -> None:
    """Remove the example paragraphs sitting between two landmark paragraphs."""
    paragraphs = doc.paragraphs
    start = end = None
    for i, p in enumerate(paragraphs):
        stripped = p.text.strip()
        if start is None and stripped == start_text:
            start = i
        elif start is not None and stripped == end_text:
            end = i
            break
    if start is None or end is None:
        return
    for p in paragraphs[start + 1 : end]:
        p._element.getparent().remove(p._element)


def _find_paragraph(doc, text: str):
    for p in doc.paragraphs:
        if p.text.strip() == text:
            return p
    return None


def _points_for(memo: ScoringMemo, attr: str) -> list:
    return getattr(memo, attr)


def build_docx(
    memo: ScoringMemo,
    research: str | None = None,
    sources: list[str] | None = None,
    template_path: Path | None = None,
) -> io.BytesIO:
    """Render the memo into the template and return it as an in-memory .docx."""
    doc = Document(str(template_path or TEMPLATE_PATH))

    # --- title ------------------------------------------------------------
    title = doc.paragraphs[0] if doc.paragraphs else None
    if title is not None:
        _set_paragraph_text(title, f"{memo.company_name} — {memo.tagline}")

    facts = memo.fact_map()
    scoring_line = memo.scoring_line()

    # --- deal facts table -------------------------------------------------
    table = doc.tables[0]
    rows = table.rows

    def put(row_index: int, column: int, value: str) -> None:
        if row_index < len(rows):
            _set_cell(rows[row_index].cells[column], value)

    for row_index, label in LEFT_ROWS:
        put(row_index, 1, facts.get(label.lower(), ""))
    for row_index, label in RIGHT_ROWS:
        put(row_index, 3, facts.get(label.lower(), ""))
    for row_index in BLANK_LEFT_ROWS:
        put(row_index, 1, "")
    for row_index in BLANK_RIGHT_ROWS:
        put(row_index, 3, "")
    put(SCORING_ROW, 3, scoring_line)

    # --- scores table -----------------------------------------------------
    scores_table = doc.tables[1]
    for column, (_, _, score, _rationale) in enumerate(memo.scored()):
        if column < len(scores_table.rows[1].cells):
            _set_cell(scores_table.rows[1].cells[column], str(score))

    # --- "Scoring: /30" heading ------------------------------------------
    for p in doc.paragraphs:
        if p.text.strip().startswith("Scoring:") and "/" in p.text:
            _set_paragraph_text(p, f"Scoring: {scoring_line}")
            break

    # --- pros / cons ------------------------------------------------------
    _delete_between(doc, "Pros:", "Cons/risks:")
    _delete_between(doc, "Cons/risks:", "Forwarding summary:")

    cons_anchor = _find_paragraph(doc, "Cons/risks:")
    if cons_anchor is not None:
        for point in _points_for(memo, "pros"):
            _insert_point_before(cons_anchor, point)

    forwarding_anchor = _find_paragraph(doc, "Forwarding summary:")
    if forwarding_anchor is not None:
        for point in _points_for(memo, "cons_risks"):
            _insert_point_before(forwarding_anchor, point)

    # --- forwarding summary list -----------------------------------------
    summary = memo.forwarding_summary
    replacements = {
        "Round terms:": f"Round terms: {summary.round_terms}".strip(),
        "Source & terms:": f"Source & terms: {summary.source_and_terms}".strip(),
        "Deadline:": f"Deadline: {summary.deadline}".strip(),
        "Scoring:": f"Scoring: {scoring_line}",
        "Dropbox:": "Dropbox:",
        "Pipedrive:": "Pipedrive:",
    }
    for p in doc.paragraphs:
        text = p.text.strip()
        for prefix, replacement in replacements.items():
            if text.startswith(prefix) and p.style.name == "List Paragraph":
                _set_paragraph_text(p, replacement)
                break

    # --- the nine detail sections ----------------------------------------
    for table_index, key in SECTION_TABLES.items():
        if table_index >= len(doc.tables):
            continue
        cell = doc.tables[table_index].rows[1].cells[0]
        _fill_cell_lines(cell, _parse_lines(getattr(memo.sections, key, "")))

    # --- appendix ---------------------------------------------------------
    _append_appendix(doc, memo, research, sources)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


def _set_paragraph_text(paragraph, text: str) -> None:
    """Replace a paragraph's text while keeping its first run's formatting."""
    for run in paragraph.runs[1:]:
        run._element.getparent().remove(run._element)
    if paragraph.runs:
        paragraph.runs[0].text = text
    else:
        paragraph.add_run(text)


def _insert_point_before(anchor, point) -> None:
    new = anchor.insert_paragraph_before("")
    run = new.add_run(f"{point.category}: ")
    run.bold = True
    new.add_run(point.point)


def _parse_lines(text: str) -> list[tuple[str, str]]:
    """Split a section's text into (label, value) pairs.

    The prompt asks for one `Label: value` line per sub-heading. Anything that
    does not follow that shape is kept as an unlabelled line rather than dropped,
    so a stray paragraph still reaches the document.
    """
    pairs: list[tuple[str, str]] = []
    for raw in (text or "").splitlines():
        line = raw.strip().lstrip("-•").strip()
        if not line:
            continue
        label, sep, value = line.partition(":")
        # A colon deep into a sentence is punctuation, not a label.
        if sep and value.strip() and len(label) <= 44 and "." not in label:
            pairs.append((label.strip(), value.strip()))
        else:
            pairs.append(("", line))
    return pairs or [("", "—")]


def _bullet(doc, text: str):
    """Add a bullet line, falling back to a plain paragraph.

    The template defines "List Paragraph" but not "List Bullet", and a template
    the fund edits later may define neither.
    """
    for style in ("List Bullet", "List Paragraph"):
        try:
            return doc.add_paragraph(text, style=style)
        except KeyError:
            continue
    return doc.add_paragraph(f"• {text}")


def _append_appendix(doc, memo: ScoringMemo, research: str | None, sources: list[str] | None) -> None:
    """Score rationales, gaps and web research, after the template's own content."""
    doc.add_paragraph()
    heading = doc.add_paragraph()
    heading.add_run("Scoring rationale").bold = True
    for _, label, score, rationale in memo.scored():
        p = doc.add_paragraph()
        p.add_run(f"{label} — {score}/3: ").bold = True
        p.add_run(rationale)

    if memo.missing_information:
        p = doc.add_paragraph()
        p.add_run("Missing information").bold = True
        for item in memo.missing_information:
            _bullet(doc, item)

    if research:
        p = doc.add_paragraph()
        p.add_run("Appendix — web research").bold = True
        note = doc.add_paragraph()
        note.add_run(
            "Gathered from public sources, not from the founders. Not verified and "
            "possibly out of date."
        ).italic = True
        for block in research.split("\n\n"):
            if block.strip():
                doc.add_paragraph(block.strip())

    if sources:
        p = doc.add_paragraph()
        p.add_run("Sources consulted").bold = True
        for source in sources:
            _bullet(doc, source)


def filename_for(memo: ScoringMemo) -> str:
    import datetime as dt

    safe = "".join(c if c.isalnum() or c in "-_ " else "" for c in memo.company_name).strip()
    safe = safe.replace(" ", "_") or "company"
    date = dt.datetime.now(dt.timezone.utc).strftime("%Y%m%d")
    return f"{safe}_scoring_memo_{date}.docx"


__all__ = ["build_docx", "filename_for", "MAX_TOTAL", "CATEGORIES"]
