"""The filled Word template: right values in the right cells, no example data left."""

from docx import Document

from vcbot.docx_memo import TEMPLATE_PATH, build_docx, filename_for
from tests.test_scoring import make_memo


def rendered(memo=None, **kwargs):
    return Document(build_docx(memo or make_memo(), **kwargs))


def all_text(doc) -> str:
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts += [c.text for c in row.cells]
    return "\n".join(parts)


def test_template_ships_with_the_package():
    assert TEMPLATE_PATH.exists()


def test_title_is_the_company_not_the_example():
    doc = rendered()
    assert "Vektor Robotics" in doc.paragraphs[0].text
    assert "Armeta" not in doc.paragraphs[0].text


def test_no_example_data_survives():
    """The template carries a worked example; none of it may leak into a memo."""
    text = all_text(rendered())
    for leftover in ["Armeta", "Artem Bosov", "armeta.ai", "Kazakhstan", "Alchemist", "26/30"]:
        assert leftover not in text, f"template example data leaked: {leftover}"


def test_deal_facts_land_in_the_right_cells():
    table = rendered().tables[0]
    assert table.rows[0].cells[0].text.strip().startswith("Round size")
    assert table.rows[0].cells[1].text.strip() == "$4M"
    assert table.rows[2].cells[3].text.strip() == "https://vektor.example"
    assert table.rows[8].cells[1].text.strip() == "$1M pre-seed"
    assert table.rows[11].cells[3].text.strip() == "$116K"


def test_fund_workflow_fields_are_left_blank():
    """VP, GP, contacts and the dates are the fund's to fill, not the bot's to invent."""
    table = rendered().tables[0]
    assert table.rows[0].cells[3].text.strip() == ""   # VP
    assert table.rows[1].cells[3].text.strip() == ""   # GP
    assert table.rows[10].cells[1].text.strip() == ""  # Contacts
    assert table.rows[11].cells[1].text.strip() == ""  # Commitment Date
    assert table.rows[12].cells[1].text.strip() == ""  # Money Transfer Date


def test_scores_row_matches_the_memo():
    memo = make_memo(market=3, tech=1, deal=0)
    row = rendered(memo).tables[1].rows[1]
    assert [c.text.strip() for c in row.cells] == ["3", "2", "2", "2", "2", "2", "2", "1", "0", "2"]


def test_scoring_total_appears_in_heading_and_facts_table():
    memo = make_memo(market=3, product=3)
    doc = rendered(memo)
    text = all_text(doc)
    assert "Scoring: 22/30" in text
    assert doc.tables[0].rows[12].cells[3].text.strip() == "22/30"


def test_pros_and_cons_are_labelled_by_category():
    text = all_text(rendered())
    assert "Traction: ARR grew 4.5x in twelve months" in text
    assert "Traction: 71% of ARR is one customer" in text


def test_the_nine_sections_are_filled():
    doc = rendered()
    text = all_text(doc)
    assert "3PL throughput" in text          # 1 Product
    assert "$4,200/arm/mo" in text           # 2 Business model
    assert "$1.4M ARR" in text               # 3 Traction
    assert "Dana Whitfield" in text          # 4 Team
    assert "Founder-led outbound" in text    # 5 Go to market
    assert "top-down" in text                # 6 Market
    assert "Locus, Fetch" in text            # 7 Competitors
    assert "4.1M pick events" in text        # 8 Technology
    assert "13x ARR" in text                 # 9 Deal/Ask


def test_section_headings_are_preserved():
    text = all_text(rendered())
    for heading in [
        "1. Product",
        "2. Business model",
        "3. Traction",
        "4. Team",
        "5. Go to market",
        "6. Market",
        "7. Competitors",
        "8. Technology",
        "9. Deal/Ask",
    ]:
        assert heading in text


def test_appendix_carries_rationales_and_research():
    text = all_text(
        rendered(research="No public footprint found.", sources=["Crunchbase — https://cb.com/x"])
    )
    assert "Scoring rationale" in text
    assert "Market — 2/3" in text
    assert "Appendix — web research" in text
    assert "No public footprint found." in text
    assert "Crunchbase — https://cb.com/x" in text


def test_appendix_omits_research_when_none_ran():
    text = all_text(rendered())
    assert "Appendix — web research" not in text
    assert "Scoring rationale" in text


def test_missing_information_is_listed():
    assert "No cohort data" in all_text(rendered())


def test_filename_is_a_docx_and_filesystem_safe():
    memo = make_memo()
    memo.company_name = "Vektor / Robotics: Inc."
    name = filename_for(memo)
    assert name.endswith(".docx")
    assert "/" not in name and ":" not in name


def test_empty_facts_do_not_crash_rendering():
    memo = make_memo()
    memo.facts = ""
    memo.rationales = ""
    for field in type(memo.sections).model_fields:
        setattr(memo.sections, field, "")
    doc = rendered(memo)
    assert doc.tables[0].rows[0].cells[1].text.strip() == ""


def test_malformed_facts_block_fills_what_it_can():
    """A model that drops or reorders a label must not misalign the table."""
    memo = make_memo()
    memo.facts = "Industry: Robotics\nnonsense line without a colon\nStage: Seed"
    table = rendered(memo).tables[0]
    assert table.rows[3].cells[3].text.strip() == "Robotics"  # Industry
    assert table.rows[4].cells[3].text.strip() == "Seed"      # Stage
    assert table.rows[0].cells[1].text.strip() == ""          # Round size, absent
